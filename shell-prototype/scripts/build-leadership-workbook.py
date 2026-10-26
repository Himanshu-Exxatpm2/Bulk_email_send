#!/usr/bin/env python3
"""Builds Bulk_Email_Use_Cases_Leadership_Review.xlsx: a short workbook for leadership to evaluate.

Reads the same reviewed set as the detailed workbook (scripts/email_use_cases.py), so the two never
disagree, but shows only what a reviewer needs: the use cases, a plain-language reason for each,
readable sample emails, and what was left out.

Run from the project root (needs openpyxl):  python3 scripts/build-leadership-workbook.py
"""
import json
import re
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))
from email_use_cases import SURFACES, UC  # noqa: E402
from email_use_cases_v1 import UC as V1  # noqa: E402
from review_v1 import REVIEW  # noqa: E402

norm = lambda t: re.sub(r"[^a-z0-9]", "", t.lower())
src = (ROOT / "src/app/data/mergeFieldCatalog.ts").read_text()
CAT = {}
for l in src.splitlines():
    if l.startswith('  {"token"'):
        f = json.loads(l.strip().rstrip(","))
        CAT.setdefault(norm(f["token"]), f)

# ── one plain-language reason per use case ───────────────────────────────────
WHY = {
    ("Site - Schedules", "Onboarding reminder"): "Onboarding gates every start. One reminder with each student's own outstanding items gets the late ones moving.",
    ("Site - Schedules", "Rotation starting soon: where to go and who to see"): "Answers the most repeated student question (where, when, who) in one message per start date.",
    ("Site - Schedules", "Preceptor heads-up for an incoming student"): "An unprepared preceptor is the commonest day-one failure. This tells them who is coming and when.",
    ("Site - Schedules", "Schedule change or cancellation notice"): "A student who misses a change becomes a no-show at a clinical site.",
    ("Site - Location Schedules", "Preceptor assignment confirmation"): "Unconfirmed assignments are why students arrive to nobody expecting them.",
    ("Site - Location Schedules", "Day-one logistics for this location"): "Only the location's coordinator can write its parking, badge and dress-code details.",
    ("Site - Location Schedules", "Invite finishing students to apply"): "Joins placements to hiring: students who already worked on the unit are the best-fit applicants.",
    ("Site - Partner Schedules", "Confirmed placements for the school"): "Schools plan staffing around confirmed placements. One roll-up per term.",
    ("Site - Partner Schedules", "Students not cleared to start"): "Onboarding is the biggest reason a student cannot start, and only the school can chase its own students.",
    ("Site - Partner Schedules", "Placement change notice to the school"): "The school must hear about a change before its students do.",
    ("Site - Location Personnel", "Upcoming students: coverage heads-up"): "Staff need advance notice of who is coming so they can plan capacity.",
    ("Site - Location Personnel", "Action needed: assign preceptors and review placements"): "Assignment is the step most often left until the last week.",
    ("Site - Location Personnel", "Requests waiting for your decision"): "168 requests were pending approval at the May baseline. Each is a school waiting on a coordinator.",
    ("Site - Personnel", "Confirm your capacity for the next block"): "No supervisor commitment means no availability, and no availability means no placement.",
    ("Site - Personnel", "Policy or compliance announcement (structured)"): "The most common staff email, in a format that says what changed, what to do and by when.",
    ("Site - Personnel", "Start-of-term responsibilities refresher"): "Roles change each term. A short prompt gets requests decided and placements assigned early.",
    ("Site - School Partners", "Contract expiring: renewal reminder"): "A lapsed contract blocks placements, and renewal is revenue.",
    ("Site - School Partners", "Status of your requests"): "Schools guess where their requests stand. One digest replaces dozens of status calls.",
    ("Site - School Partners", "Confirm the program contacts on file"): "Every email to a school depends on its contacts being right.",
    ("Site - Partner Contacts", "Request needs more information"): "Most stalled requests are waiting on one missing answer.",
    ("Site - Partner Contacts", "Request approved: next steps"): "Every approval is followed by the school scheduling students.",
    ("Site - Partner Contacts", "Request not approved: reason and next steps"): "More than half of requests are not approved. How a site says no decides whether the school comes back.",
    ("Site - Jobs Candidates", "Application received"): "Every applicant should hear back at once.",
    ("Site - Jobs Candidates", "Interview invitation"): "Interviews run in rounds, and a missed invitation loses a hire.",
    ("Site - Jobs Candidates", "Application status update"): "Most applicants are not hired. Silence damages the employer brand.",
    ("Site - Jobs Candidates", "Offer expiring soon: reminder"): "A lapsed offer is a lost hire the employer already paid to recruit.",
    ("School - Schedules", "Remind students to onboard"): "Onboarding is the school's main blocker and its most repeated nudge.",
    ("School - Schedules", "Schedule change or cancellation notice"): "Students should hear about a changed placement from their school, not by turning up.",
    ("School - Schedules", "Placement confirmed: what to expect"): "Every confirmed student needs the same details in one message.",
    ("School - Schedules", "Brief your clinical instructors"): "Instructors are the school's link on site but have no single view of who is blocked.",
    ("School - Requests", "Follow up on a pending request"): "Schools wait on decisions that gate their whole term.",
    ("School - Requests", "Responding to a request for more information"): "A request stalls until the school answers.",
    ("School - Requests", "Withdraw requests we no longer need"): "Frees slots for other schools and clears dead requests from the backlog.",
    ("School - Requests", "Approved: thank you and next steps"): "An approval is only half done until the students are named.",
}
for u in UC:
    assert (u["surface"], u["name"]) in WHY, (u["surface"], u["name"])

# Short "how we will know it worked" lines, reused from the detailed build.
METRIC = {
    "Onboarding reminder": "Share of recipients cleared within 7 days of the send.",
    "Rotation starting soon: where to go and who to see": "Fewer 'where / who' questions to coordinators before the start.",
    "Preceptor heads-up for an incoming student": "Share of students with an informed preceptor on day one.",
    "Schedule change or cancellation notice": "No-shows among affected students.",
    "Preceptor assignment confirmation": "Share of preceptors confirmed 5 days before the block.",
    "Day-one logistics for this location": "Day-one 'where do I go' calls to the location.",
    "Invite finishing students to apply": "Applications and first hires from the cohort.",
    "Confirmed placements for the school": "School queries asking for a placement list.",
    "Students not cleared to start": "Share of listed students cleared a week later.",
    "Placement change notice to the school": "Schools learning of a change from the student first.",
    "Upcoming students: coverage heads-up": "Coverage conflicts raised after the block starts.",
    "Action needed: assign preceptors and review placements": "Placements with no preceptor 7 days before start.",
    "Requests waiting for your decision": "Median days from request to decision; size of the pending backlog.",
    "Confirm your capacity for the next block": "Share of staff who reply before availabilities are posted.",
    "Policy or compliance announcement (structured)": "Share of staff completing the action by the date.",
    "Start-of-term responsibilities refresher": "Requests decided and placements assigned in the first two weeks.",
    "Contract expiring: renewal reminder": "Contracts renewed before expiry.",
    "Status of your requests": "Fewer status-chasing emails and calls from schools.",
    "Confirm the program contacts on file": "Contacts corrected; schools with no contact.",
    "Request needs more information": "Days from 'info requested' to the school's reply.",
    "Request approved: next steps": "Days from approval to students being scheduled.",
    "Request not approved: reason and next steps": "Share of declined schools that request again next term.",
    "Application received": "Applicants asking 'did you get my application'.",
    "Interview invitation": "Interview confirmations within 48 hours.",
    "Application status update": "Applicant complaints; reapplication rate.",
    "Offer expiring soon: reminder": "Offers answered before expiry.",
    "Remind students to onboard": "Share of recipients cleared within 7 days of the send.",
    "Schedule change or cancellation notice": "No-shows among affected students.",
    "Placement confirmed: what to expect": "Students arriving prepared; onboarding done before the deadline.",
    "Brief your clinical instructors": "Instructors contacting not-cleared students before the start.",
    "Follow up on a pending request": "Days to decision on the requests chased.",
    "Responding to a request for more information": "Days from info request to reply.",
    "Withdraw requests we no longer need": "Slots released early.",
    "Approved: thank you and next steps": "Days from approval to student details submitted.",
}

# Reasons an idea was left out, one line each.
NOT_INCLUDED = [
    ("Site - Location Schedules", "Briefing for the people at the location", "Same email as the coverage heads-up on Location Personnel."),
    ("Site - Location Schedules", "Location notice to placed students", "Rare, and the schedule change notice already reaches the same students."),
    ("Site - Partner Schedules", "Students who have not signed in", "Resend Invite already handles sign-in. Widened into 'Students not cleared to start'."),
    ("Site - Partner Schedules", "Clinical instructor roster coordination", "Tells instructors what their own school already knows."),
    ("Site - Location Personnel", "Announcement or policy update", "A blank message box is not a template. Kept as a structured version on the Personnel screen."),
    ("Site - Location Personnel", "Sign in to Exxat One", "Resend Invite already covers it."),
    ("Site - Personnel", "Sign-in nudge for staff who have not logged in", "Resend Invite already covers it."),
    ("Site - Personnel", "Confirm your details and locations", "Better as a profile-review flow in the product than as a bulk email."),
    ("Site - School Partners", "Open requests waiting on the school", "Wrongly framed: pending requests wait on the site. Replaced by a status digest for schools."),
    ("Site - School Partners", "Partner announcement or term update", "A blank message box, and a System Announcement template already exists."),
    ("Site - Partner Contacts", "Contract renewal for this school", "Duplicate of the renewal reminder on School Partners."),
    ("Site - Partner Contacts", "Placement activity check-in", "Vague, with no clear moment to send it."),
    ("Site - Jobs Candidates", "Offer", "Offers carry individual terms. Replaced by an offer-expiring reminder."),
    ("School - Schedules", "First sign-in nudge", "Resend Invite already covers it."),
    ("School - Schedules", "Heads-up to the site about incoming students", "The site already sees its schedules. Replaced by briefing the school's instructors."),
    ("School - Requests", "Request to change dates", "Negotiated one request at a time. Replaced by withdrawing requests no longer needed."),
]
assert len(NOT_INCLUDED) == len([k for k, v in REVIEW.items() if v[1] == "Replace"]) + 1, len(NOT_INCLUDED)

# ── sample data so emails read like emails ───────────────────────────────────
LISTS = {
    "outstandingItems": ["Immunization records", "Background check", "CPR certification"],
    "studentSchedules": ["Aaliyah Washington — MedStar Washington Hospital Center – ICU — Jan 12 – Apr 2, 2027",
                         "Marcus Okafor — MedStar Washington Hospital Center – ICU — Jan 12 – Apr 2, 2027",
                         "Destiny Robinson — MedStar Washington Hospital Center – ED — Jan 19 – Apr 9, 2027"],
    "studentsNotCleared": ["Marcus Okafor — MedStar Washington Hospital Center – ICU — starts January 12, 2027 — Not Started",
                           "Destiny Robinson — MedStar Washington Hospital Center – ED — starts January 19, 2027 — Pending"],
    "openRequests": ["Request 3558157000 – MedStar ICU – Spring 2027 – Pending",
                     "Request 4981028000 – MedStar ED – Spring 2027 – Info Requested"],
    "requestStatuses": ["Request 3558157000 – MedStar ICU – Spring 2027 – Pending",
                        "Request 4981028000 – MedStar ED – Spring 2027 – Info Requested",
                        "Request 7799145000 – MedStar Med/Surg – Fall 2026 – Approved"],
    "programContacts": ["Dr. Renee Blackwell — r.blackwell@bowiestate.edu", "Marcus Ellery — m.ellery@bowiestate.edu"],
}
COMMON = dict(
    rotationName="MedStar ICU – Spring 2027", assignedScheduleRef="MedStar ICU – Spring 2027",
    startDate="January 12, 2027", endDate="April 2, 2027", scheduleDateRange="Jan 12 – Apr 2, 2027",
    onboardingDeadline="December 29, 2026", today="October 22, 2026", submittedDate="October 13, 2026",
    requestId="3558157000", requestedSlots="2", approvedSlotCount="2", programType="BSN", slotStatus="Approved",
    requestedLocation="MedStar Washington Hospital Center – ICU", contractExpirationDate="March 31, 2027",
    daysToContractExpiry="160", applicationDate="October 14, 2026", schoolName="Bowie State University",
    siteName="MedStar Health", hospitalName="MedStar Health", senderName="Kierra Mango", scheduleStatus="Upcoming",
    endDateStudent="April 2, 2027", pendingStudents="5", scheduleCount="229",
)
PERSONA = {
    "Student": dict(firstName="Aaliyah", lastName="Washington", recipientName="Aaliyah Washington", studentName="Aaliyah Washington", email="a.washington@bowiestate.edu"),
    "School Contact": dict(firstName="Renee", lastName="Blackwell", recipientName="Dr. Renee Blackwell", email="r.blackwell@bowiestate.edu", schoolContactName="Dr. Renee Blackwell"),
    "Staff-Personnel": dict(firstName="Sandra", lastName="Okonkwo", recipientName="Sandra Okonkwo", staffName="Sandra Okonkwo", email="s.okonkwo@medstarhealth.org"),
    "Preceptor": dict(firstName="David", lastName="Whitfield", recipientName="David Whitfield", preceptorName="David Whitfield", email="d.whitfield@medstarhealth.org"),
    "Placement Faculty/CI": dict(firstName="Denise", lastName="Carrington", recipientName="Denise Carrington", facultyName="Denise Carrington", email="d.carrington@bowiestate.edu"),
    "Candidate": dict(firstName="Priya", lastName="Raman", recipientName="Priya Raman", candidateName="Priya Raman", email="p.raman@example.com"),
}
SENDER_ORG = {"site": "MedStar Health Clinical Education", "school": "Bowie State University Clinical Placement Office"}
SENDER_FILLS = {"changeDescription": "[the sender describes the change here]", "messageBody": "[the sender writes the message here]",
                "interviewDate": "[interview date]", "interviewTime": "[interview time]", "interviewFormat": "[interview format]",
                "interviewerName": "[interviewer name]", "interviewerTitle": "[interviewer title]", "offerExpirationDate": "[offer expiry date]",
                "recruiterName": "[recruiter name]", "confirmationDeadline": "[confirmation deadline]", "newStartDate": "[new start date]", "newEndDate": "[new end date]"}


def sample(text, u):
    portal = SURFACES[u["surface"]]["portal"]
    values = dict(COMMON)
    values.update(PERSONA[u["persona"]])
    if u["surface"] == "School - Requests":
        values.update(firstName="Alex", recipientName="Alex Rivera")  # the site contact for the request
    if u["persona"] == "Candidate":
        values.update(jobTitle="ICU Registered Nurse", jobLocation="Washington, DC", department="Critical Care", stage="Under Review",
                      compensation="$78,000 – 96,000", jobsPortalUrl="https://one.exxat.com/jobs", hospitalName="MedStar Health")
    values["senderOrganization"] = SENDER_ORG[portal]

    def each(m):
        items = LISTS.get(norm(m.group(1)) and m.group(1), [])
        return "".join("• " + i + "\n" for i in items).rstrip("\n")

    out = re.sub(r"\{\{#each\s+([^}\s]+)\s*\}\}.*?\{\{/each\}\}", each, text, flags=re.S)

    def one(m):
        t = m.group(1).strip()
        if t in SENDER_FILLS:
            return SENDER_FILLS[t]
        if t in values:
            return values[t]
        for k, v in values.items():
            if norm(k) == norm(t):
                return v
        f = CAT.get(norm(t))
        if f and f["example"]:
            return f["example"]
        return "[" + (f["label"] if f else t) + "]"

    return re.sub(r"\{\{\s*([^#/}][^}]*?)\s*\}\}", one, out)


# ── priority and wording helpers ─────────────────────────────────────────────
def priority(score):
    return "High" if score >= 22 else ("Medium" if score >= 20 else "Standard")


def builds_on(maps_to):
    if maps_to.startswith("Net new"):
        return "New"
    m = re.match(r"^([^(.]+)", maps_to)
    return "Builds on: " + m.group(1).strip()


FRIENDLY = {"Student": "Students", "School Contact": "School contacts", "Staff-Personnel": "Site staff", "Preceptor": "Preceptors",
            "Placement Faculty/CI": "Clinical instructors", "Candidate": "Candidates"}
SHORT_SCREEN = {k: v["screen"].replace(" (not wired in this prototype)", "") for k, v in SURFACES.items()}

# ── workbook ─────────────────────────────────────────────────────────────────
from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

NAVY = "1F3A8A"
HEAD = PatternFill("solid", fgColor=NAVY)
BAND = PatternFill("solid", fgColor="EEF2FF")
thin = Side(style="thin", color="D1D5DB")
BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
n = len(UC)
screens = Counter(u["surface"] for u in UC)
new_count = sum(1 for u in UC if builds_on(u["maps_to"]) == "New")
prio = Counter(priority(u["score"]) for u in UC)

# 1 Overview
ws = wb.active
ws.title = "Overview"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 46
ws.column_dimensions["C"].width = 100
ws["B1"] = "Bulk Email: recommended use cases"
ws["B1"].font = Font(bold=True, size=18, color=NAVY)
ws["B2"] = "For leadership review and validation"
ws["B2"].font = Font(size=11, color="6B7280")


def para(row, label, text, bold_label=True):
    ws.cell(row=row, column=2, value=label).font = Font(bold=bold_label, color=NAVY if bold_label else "000000")
    c = ws.cell(row=row, column=3, value=text)
    c.alignment = WRAP
    ws.cell(row=row, column=2).alignment = WRAP


r = 4
para(r, "What this is", f"The best {n} email use cases for the 10 places Bulk Email appears in Exxat One, with a ready-to-use email for each. Each one answers a recurring moment in the placement journey where the same message needs to reach many people at once."); r += 1
para(r, "What we need from you", "Open the Use Cases sheet and mark each row Approve, Approve with changes, or Drop in the Decision column, with a note where you want a change. Then answer the three questions below."); r += 2

ws.cell(row=r, column=2, value="At a glance").font = Font(bold=True, size=12, color=NAVY); r += 1
for label, val in [
    ("Use cases", str(n)),
    ("Screens covered", "10: 8 on the Site portal and 2 on the School portal."),
    ("Priority", f"{prio['High']} High, {prio['Medium']} Medium, {prio['Standard']} Standard (judgement scores out of 25; High is 22 or more)"),
    ("Build on an existing template", f"{n - new_count} of {n}. The other {new_count} are new."),
    ("Left out on purpose", f"{len(NOT_INCLUDED)} ideas, with the reason for each, on the Not Included sheet."),
]:
    para(r, label, val, bold_label=False); r += 1
r += 1

ws.cell(row=r, column=2, value="How the use cases were chosen").font = Font(bold=True, size=12, color=NAVY); r += 1
for t in [
    "Started from the placement journey (onboarding, before a rotation starts, assignment, change, requests, contracts, hiring) and kept only the moments that repeat every term for many people.",
    "Scored each candidate out of 25 on five tests: how many people and how often, what goes wrong if it is not sent, whether it is one message to many, whether it is distinct from another screen, and whether the data to fill it in exists.",
    "Dropped anything the portal already does natively (for example Resend Invite for login), anything that needs individual terms (such as offers), and anything that duplicated an email on another screen.",
    "Every use case replaced is stronger than the one it replaced; the Not Included sheet shows each swap.",
]:
    c = ws.cell(row=r, column=3, value="• " + t)
    c.alignment = WRAP
    r += 1
r += 1

ws.cell(row=r, column=2, value="Questions for leadership").font = Font(bold=True, size=12, color=NAVY); r += 1
for q, t in [
    ("1. Measure it", "No send data exists yet, so priorities are judgement. Should Bulk Email record which use case every send belongs to, so we can replace judgement with usage?"),
    ("2. Jobs", "The Jobs use cases assume hiring volume. The one Site tenant we checked showed no jobs. Is that representative before we invest here?"),
    ("3. Faculty assignment", "Faculty assignment is fully manual today. The strongest fix is asking schools to assign clinical instructors, which needs instructor data per schedule that does not exist yet. Do we want it built?"),
]:
    para(r, q, t, bold_label=True); r += 1
r += 1
ws.cell(row=r, column=2, value="Caveats").font = Font(bold=True, size=12, color=NAVY); r += 1
for t in [
    "Scores and priorities are judgement, not measured data.",
    "Some baseline figures (168 requests pending approval, 47% approval rate, 196 open availabilities) are the May 2026 baseline.",
    "The Jobs screen's Bulk Email button is not built in the prototype yet; those use cases are for when it is.",
    "A separate detailed workbook holds every merge field, the full scoring and the paste-ready templates for the team that builds this.",
]:
    ws.cell(row=r, column=3, value="• " + t).alignment = WRAP
    r += 1

# 2 Use Cases
ws = wb.create_sheet("Use Cases")
H = ["#", "Screen", "Use case", "Who receives it", "When it is sent", "Why it matters", "Based on", "Priority", "How we will know it worked", "Decision", "Comments"]
ws.append(H)
for i, h in enumerate(H, 1):
    c = ws.cell(row=1, column=i)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = HEAD
    c.alignment = Alignment(wrap_text=True, vertical="center")
    c.border = BORDER
last_screen, band = None, False
for i, u in enumerate(UC, 1):
    if u["surface"] != last_screen:
        band = not band
        last_screen = u["surface"]
    ws.append([i, f'{SURFACES[u["surface"]]["portal"].capitalize()}: {SHORT_SCREEN[u["surface"]]}', u["name"], u["recipients"], u["when"],
               WHY[(u["surface"], u["name"])], builds_on(u["maps_to"]), priority(u["score"]), METRIC[u["name"]], None, None])
    for c in ws[ws.max_row]:
        c.alignment = WRAP
        c.border = BORDER
        if band:
            c.fill = BAND
    ws.cell(row=ws.max_row, column=3).font = Font(bold=True)
widths = [4, 30, 38, 34, 32, 52, 28, 10, 40, 18, 36]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(H))}{ws.max_row}"
dv = DataValidation(type="list", formula1='"Approve,Approve with changes,Drop"', allow_blank=True)
dv.error = "Choose Approve, Approve with changes, or Drop"
ws.add_data_validation(dv)
dv.add(f"J2:J{ws.max_row}")
rng = f"J2:J{ws.max_row}"
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Approve"'], fill=PatternFill("solid", bgColor="BBF7D0")))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Approve with changes"'], fill=PatternFill("solid", bgColor="FDE68A")))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Drop"'], fill=PatternFill("solid", bgColor="FECACA")))
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# 3 Sample Emails
ws = wb.create_sheet("Sample Emails")
H = ["#", "Use case", "Screen", "Subject", "Email (with example data)"]
ws.append(H)
for i, h in enumerate(H, 1):
    c = ws.cell(row=1, column=i)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = HEAD
    c.alignment = Alignment(wrap_text=True, vertical="center")
    c.border = BORDER
for i, u in enumerate(UC, 1):
    ws.append([i, u["name"], f'{SURFACES[u["surface"]]["portal"].capitalize()}: {SHORT_SCREEN[u["surface"]]}', sample(u["subject"], u), sample(u["body"].replace("<br>", ""), u)])
    for c in ws[ws.max_row]:
        c.alignment = WRAP
        c.border = BORDER
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=4).font = Font(bold=True)
for i, w in enumerate([4, 34, 30, 48, 100], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:E{ws.max_row}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# 4 Not Included
ws = wb.create_sheet("Not Included")
H = ["Screen", "Idea we dropped", "Why"]
ws.append(H)
for i, h in enumerate(H, 1):
    c = ws.cell(row=1, column=i)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = HEAD
    c.border = BORDER
for surface, name, why in NOT_INCLUDED:
    ws.append([f'{SURFACES[surface]["portal"].capitalize()}: {SHORT_SCREEN[surface]}', name, why])
    for c in ws[ws.max_row]:
        c.alignment = WRAP
        c.border = BORDER
for i, w in enumerate([36, 46, 90], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"

out = ROOT / "Bulk_Email_Use_Cases_Leadership_Review.xlsx"
wb.save(out)
print("workbook ->", out.name, f"| {n} use cases, {prio}")
