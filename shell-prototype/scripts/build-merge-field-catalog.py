#!/usr/bin/env python3
"""Builds the merge-field catalogue from scripts/merge_field_entries.py.

Writes:
  src/app/data/mergeFieldCatalog.ts   what the compose panel, both template pages and the engine read
  Merge_Fields_Catalog.xlsx           the workbook for the team

Run from the project root (needs openpyxl):  python3 scripts/build-merge-field-catalog.py
"""
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))
from merge_field_entries import ENTRIES, REMOVED  # noqa: E402

AUD = {"S": "Student", "C": "School Contact", "P": "Partner Contact", "F": "Staff-Personnel",
       "J": "Candidate", "R": "Preceptor", "L": "Placement Faculty/CI"}
AUD_ORDER = ["Student", "School Contact", "Partner Contact", "Staff-Personnel", "Candidate", "Preceptor", "Placement Faculty/CI"]
PORTAL = {"b": ["site", "school"], "s": ["site"], "c": ["school"]}
GROUP_ORDER = ["Recipient", "Sender & Organization", "Student", "Schedule", "Location", "Availability", "Slot Request",
               "School / Partner", "Contract", "Site Personnel", "Preceptor", "Placement Faculty / CI",
               "Candidate & Job", "Schedules List", "Group Placement", "Links", "Written by Sender"]

# Which sections of each portal a group is useful in.
SITE_SECTIONS = ["Slot Requests", "Schedules", "Onboarding", "Jobs", "General"]
SCHOOL_SECTIONS = ["Requests", "Schedules & Onboarding", "General"]
USED_IN = {
    "Recipient": (SITE_SECTIONS, SCHOOL_SECTIONS),
    "Sender & Organization": (SITE_SECTIONS, SCHOOL_SECTIONS),
    "Student": (["Schedules", "Onboarding", "General"], ["Schedules & Onboarding", "General"]),
    "Schedule": (["Schedules", "Onboarding", "General"], ["Schedules & Onboarding", "General"]),
    "Location": (["Slot Requests", "Schedules", "Onboarding", "General"], ["Requests", "Schedules & Onboarding", "General"]),
    "Availability": (["Slot Requests", "General"], ["Requests", "General"]),
    "Slot Request": (["Slot Requests", "General"], ["Requests", "General"]),
    "School / Partner": (["Slot Requests", "Schedules", "Onboarding", "General"], ["Requests", "Schedules & Onboarding", "General"]),
    "Group Placement": (["Schedules", "Onboarding", "General"], ["Schedules & Onboarding", "General"]),
    "Contract": (["Slot Requests", "General"], []),
    "Site Personnel": (["Schedules", "General"], []),
    "Preceptor": (["Schedules", "Onboarding", "General"], []),
    "Placement Faculty / CI": (["Schedules", "General"], ["Schedules & Onboarding", "General"]),
    "Candidate & Job": (["Jobs"], []),
    "Schedules List": (["Schedules", "Onboarding", "General"], ["Schedules & Onboarding", "General"]),
    "Links": (SITE_SECTIONS, SCHOOL_SECTIONS),
    "Written by Sender": (SITE_SECTIONS, SCHOOL_SECTIONS),
}

norm = lambda t: re.sub(r"[^a-z0-9]", "", t.lower())


def audiences_of(e):
    return [AUD[c] for c in e["aud"]]


# ── Pull in anything the three old tables had that the new table does not ─────
LEGACY_JSON = ROOT / "scripts" / "legacy_merge_fields.json"


def legacy_pairs():
    """The tokens the three pre-catalogue tables held, as (audience, label, token).
    Read once from those files and saved, so they survive the tables being replaced."""
    if LEGACY_JSON.exists():
        return [tuple(x) for x in json.loads(LEGACY_JSON.read_text())]
    files = [ROOT / "src/app/components/email/ComposeEmailWizard.tsx",
             ROOT / "src/app/components/email/EmailConfigurationPage.tsx",
             ROOT / "src/app/components/school/SchoolConfigurationPage.tsx"]
    out = []
    for f in files:
        text = f.read_text()
        m = re.search(r"const MERGE_FIELDS[^=]*=\s*\{(.*?)\n\};", text, re.S)
        if not m:
            continue
        for am in re.finditer(r"(?:'([^']+)'|(\w+)):\s*\[(.*?)\],", m.group(1), re.S):
            aud = am.group(1) or am.group(2)
            for lm in re.finditer(r"label:\s*'([^']*)',\s*token:\s*'\{\{([^}]+)\}\}'", am.group(3)):
                out.append((aud, lm.group(1), lm.group(2)))
    LEGACY_JSON.write_text(json.dumps(out, indent=1))
    return out


known = {norm(e["token"]) for e in ENTRIES} | {norm(t) for t in REMOVED}  # removed fields must not return via the old tables
extra = []
for aud, label, token in legacy_pairs():
    if norm(token) not in known:
        extra.append(dict(token=token, label=label, group="Written by Sender", aud={v: k for k, v in AUD.items()}[aud],
                          portal="b", prod="Typed by the sender when composing", example="", how="author",
                          kind="text", desc="Written by the sender at send time; not filled from a record.", insert=None))
        known.add(norm(token))
ENTRIES.extend(extra)

# Merge duplicate tokens that differ only by audience letters.
by_tok = {}
for e in ENTRIES:
    k = norm(e["token"])
    if k in by_tok:
        by_tok[k]["aud"] = "".join(sorted(set(by_tok[k]["aud"]) | set(e["aud"])))
    else:
        by_tok[k] = e
ENTRIES[:] = list(by_tok.values())
ENTRIES.sort(key=lambda e: GROUP_ORDER.index(e["group"]))  # stable: keeps authoring order inside a group

# ── TypeScript ───────────────────────────────────────────────────────────────
def ts_entry(e):
    site, school = USED_IN[e["group"]]
    portals = PORTAL[e["portal"]]
    used = (site if "site" in portals else []) + (school if "school" in portals else [])
    d = dict(token=e["token"], label=e["label"], group=e["group"], description=e["desc"],
             audiences=audiences_of(e), portals=portals, kind=e["kind"], example=e["example"],
             source=e["how"], prod=e["prod"])
    if e["insert"]:
        d["insert"] = e["insert"]
    return "  " + json.dumps(d, ensure_ascii=False) + ","


ts = """// GENERATED by scripts/build-merge-field-catalog.py from scripts/merge_field_entries.py.
// Edit the Python table and re-run it; do not edit this file by hand.
//
// One list of every merge field, read by: the compose panel's Insert Merge Field menu,
// the Email Setup template editor (site and school) and the merge engine (messageMerge.ts).
import type { AudienceType } from './mockEmailTemplates';

export type MergeFieldSource = 'record' | 'derived' | 'sample' | 'tenant' | 'author';
export type MergeFieldKind = 'text' | 'date' | 'number' | 'link' | 'list' | 'table';
export type Portal = 'site' | 'school';

export interface MergeFieldDef {
  /** Token name; spelling-agnostic at render time ({{first_name}} and {{firstName}} are the same). */
  token: string;
  label: string;
  group: string;
  description: string;
  audiences: AudienceType[];
  portals: Portal[];
  kind: MergeFieldKind;
  /** Sample value, also used to fill 'sample' fields the prototype has no data for. */
  example: string;
  /** record = recipient's own record; derived = computed; sample = prototype placeholder; tenant = constant; author = typed by sender. */
  source: MergeFieldSource;
  /** Where the value lives in the live product. */
  prod: string;
  /** What the menu inserts, when it is more than the bare token (list blocks). */
  insert?: string;
}

export const MERGE_FIELD_GROUPS: string[] = %s;

export const MERGE_FIELD_CATALOG: MergeFieldDef[] = [
%s
];

const norm = (t: string) => t.toLowerCase().replace(/[^a-z0-9]/g, '');

/** Fields available to one audience on one portal, in group order. */
export function mergeFieldsFor(audience: AudienceType, portal?: Portal): MergeFieldDef[] {
  return MERGE_FIELD_CATALOG.filter((f) => f.audiences.includes(audience) && (!portal || f.portals.includes(portal)));
}

/** Union across several audiences (a send to more than one persona), de-duplicated by token. */
export function mergeFieldsForAudiences(audiences: AudienceType[], portal?: Portal): MergeFieldDef[] {
  const seen = new Set<string>();
  return MERGE_FIELD_CATALOG.filter((f) => {
    if (!f.audiences.some((a) => audiences.includes(a)) || (portal && !f.portals.includes(portal))) return false;
    if (seen.has(norm(f.token))) return false;
    seen.add(norm(f.token));
    return true;
  });
}

export const mergeToken = (f: MergeFieldDef) => f.insert ?? `{{${f.token}}}`;
""" % (json.dumps(GROUP_ORDER), "\n".join(ts_entry(e) for e in ENTRIES))
(ROOT / "src/app/data/mergeFieldCatalog.ts").write_text(ts)
print("catalogue:", len(ENTRIES), "fields ->", "src/app/data/mergeFieldCatalog.ts")

# ── Excel ────────────────────────────────────────────────────────────────────
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

wb = Workbook()
HEAD = PatternFill("solid", fgColor="1F3A8A")
BAND = PatternFill("solid", fgColor="EEF2FF")
thin = Side(style="thin", color="D1D5DB")
BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(ws, headers, rows, widths):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
            c.border = BORDER
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 28


HOW = {"record": "Record", "derived": "Derived", "sample": "Sample (prototype)", "tenant": "Sender / platform", "author": "Written by sender"}
legacy_tokens = {norm(t) for _, _, t in legacy_pairs()}

# 1 Summary
ws = wb.active
ws.title = "Summary"
ws["A1"] = "Exxat One: merge field catalogue"
ws["A1"].font = Font(bold=True, size=16, color="1F3A8A")
lines = [
    "The merge fields a user can drop into a Bulk Email, grouped by what they describe and by which recipient they apply to. Reviewed screen by screen: only fields a real email would use are kept.",
    "Built from a walk through the live Site and School portals (Locations, Personnel, School Partners, Contracts, Availability, Slot Requests, Schedules, Jobs; School Requests and Schedules) plus the prototype's own records.",
    "",
    "How to read it",
    "• Where Bulk Email Appears: one sheet for each place Bulk Email is used (10 screens). Each sheet lists exactly the fields that screen can offer, with description, example and where the value lives.\n• All Merge Fields: the master list. 'Audiences' says which recipient types can use the field; 'Used in' says which Bulk Email sections it suits.",
    "• By Persona: the same list as a grid, one column per recipient type.",
    "• Removed Fields: every field cut in the review pass and why (duplicate, internal value, the recipient's own data, sensitive identifier, or no email use case).",
        "• Portal Survey: what was found on each live screen and which merge field it became.",
    "• Lists: blocks that repeat one line per item (a school's program contacts, a preceptor's students).",
    "",
    "Value source column",
    "• Record: read from the recipient's own record, or a record it links to.",
    "• Derived: computed (dates, counts, durations).",
    "• Sender / platform: the sender's name, organisation, links, today's date.",
    "• Sample (prototype): the live product has this field but the prototype has no data for it, so it fills with the example shown. Production fills it from the 'Where it lives' column.",
    "• Written by sender: typed at send time; never auto-filled.",
    "",
    "Behaviour carried over from the prototype",
    "• A field with no value for a recipient is never sent blank: that recipient is held back and flagged on the Preview step.",
    "• Values are HTML-escaped when merged into a body. The live data contains records with HTML and script text in their names (for example a personnel record named '\"/><h1>hello</h1> test'), so this matters.",
    "• Token spelling does not matter: {{first_name}}, {{firstName}} and {{First Name}} are the same field.",
]
for i, t in enumerate(lines, 3):
    ws.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical="top")
    if t in ("How to read it", "Value source column", "Behaviour carried over from the prototype"):
        ws.cell(row=i, column=1).font = Font(bold=True)
ws.column_dimensions["A"].width = 130
r0 = len(lines) + 5
ws.cell(row=r0, column=1, value="Counts by group").font = Font(bold=True)
for j, g in enumerate(GROUP_ORDER, r0 + 1):
    n = sum(1 for e in ENTRIES if e["group"] == g)
    ws.cell(row=j, column=1, value=f"{g}: {n}")
ws.cell(row=r0 + len(GROUP_ORDER) + 1, column=1, value=f"Total: {len(ENTRIES)} merge fields").font = Font(bold=True)

# 2 Merge Fields
ws = wb.create_sheet("All Merge Fields")
rows = []
for e in ENTRIES:
    site, school = USED_IN[e["group"]]
    portals = PORTAL[e["portal"]]
    used = ([f"Site: {s}" for s in site] if "site" in portals else []) + ([f"School: {s}" for s in school] if "school" in portals else [])
    rows.append([
        "{{%s}}" % e["token"], e["label"], e["group"], e["desc"], ", ".join(audiences_of(e)),
        " + ".join(p.capitalize() for p in portals), "; ".join(used), e["prod"],
        e["example"] or "(typed by sender)", e["kind"].capitalize(), HOW[e["how"]],
        "Existing" if norm(e["token"]) in legacy_tokens or e["token"] in ("firstName", "lastName", "email") else "New",
    ])
sheet(ws, ["Token", "Label", "Group", "What it is", "Audiences", "Portal", "Used in", "Where it lives (live product)",
           "Example", "Type", "Value source", "Status"], rows, [30, 30, 22, 46, 34, 14, 46, 46, 34, 10, 18, 10])
for row in ws.iter_rows(min_row=2):
    if row[0].row % 2 == 0:
        for c in row:
            c.fill = BAND

# 3 By Persona
ws = wb.create_sheet("By Persona")
rows = []
for e in ENTRIES:
    rows.append([e["group"], "{{%s}}" % e["token"], e["label"]] + ["✓" if c in e["aud"] else "" for c in "SCPFJRL"])
sheet(ws, ["Group", "Token", "Label"] + AUD_ORDER, rows, [22, 30, 30] + [14] * 7)
for row in ws.iter_rows(min_row=2, min_col=4):
    for c in row:
        c.alignment = Alignment(horizontal="center", vertical="top")
ws.append([])
tot = ["", "", "Fields available"]
for c in "SCPFJRL":
    tot.append(sum(1 for e in ENTRIES if c in e["aud"]))
ws.append(tot)
for c in ws[ws.max_row]:
    c.font = Font(bold=True)

# 4 One sheet per place Bulk Email appears
SCHED_P = ["Student", "Preceptor", "Staff-Personnel", "School Contact", "Placement Faculty/CI"]
SURFACES = [
    # sheet name, portal, screen, where, personas, note
    ("Site - Schedules", "site", "Schedules (Group and Individual views)", "Schedules > Schedules List", SCHED_P,
     "Bulk Email is hidden under All Schedules. One compose panel; tick one or more personas and their recipients are added from the selected schedules."),
    ("Site - Location Schedules", "site", "Location > Schedules tab", "Locations > a location > Schedules", SCHED_P,
     "Same list and personas as the Schedules page, scoped to the location."),
    ("Site - Partner Schedules", "site", "School Partner > Schedules tab", "School Partners > a school > Schedules", SCHED_P,
     "Same list and personas as the Schedules page, scoped to the school."),
    ("Site - Location Personnel", "site", "Location > Personnel tab", "Locations > a location > Personnel", ["Staff-Personnel"],
     "Compose panel has a person picker with role chips; nothing is pre-selected."),
    ("Site - Personnel", "site", "Personnel (list)", "Site Details > Personnel", ["Staff-Personnel"],
     "Tick staff, then Bulk Email > Compose Email."),
    ("Site - School Partners", "site", "School Partners (list)", "Site Details > School Partners", ["School Contact"],
     "Every program contact of each ticked school is added."),
    ("Site - Partner Contacts", "site", "School Partner > About > Program Contacts", "School Partners > a school > About", ["School Contact"],
     "Tick contacts in the Program Contacts table."),
    ("Site - Jobs Candidates", "site", "Jobs > Candidates (not wired in this prototype)", "Jobs > a job > Applicants", ["Candidate"],
     "Site template section 'Jobs' already exists; the Bulk Email button on the candidate list is not built in this prototype yet."),
    ("School - Schedules", "school", "Schedules & Onboarding", "School portal > Schedules", ["Student", "Staff-Personnel", "Placement Faculty/CI", "School Contact"],
     "One compose panel with personas: Students, Site Personnel, Placement Faculty / CI, School Program Contacts."),
    ("School - Requests", "school", "Requests", "School portal > Requests", ["School Contact"],
     "Goes to the site contact for the request. Fallback when no contact is on file is a product rule, not modelled here."),
]
FIELD_HEADS = ["Token", "Label", "Group", "What it is", "Available for", "Where it lives (live product)", "Example", "Type", "Value source", "Status"]
FIELD_W = [30, 32, 22, 50, 38, 48, 36, 10, 18, 10]
SECTION_BY_SURFACE = {"Site - Schedules": "Schedules, Onboarding", "Site - Location Schedules": "Schedules, Onboarding", "Site - Partner Schedules": "Schedules, Onboarding",
                      "Site - Location Personnel": "General", "Site - Personnel": "General", "Site - School Partners": "General",
                      "Site - Partner Contacts": "General", "Site - Jobs Candidates": "Jobs",
                      "School - Schedules": "Schedules & Onboarding", "School - Requests": "Requests"}
index_rows = []
surface_sheets = []
for name, portal, screen, where, personas, note in SURFACES:
    ws = wb.create_sheet(name)
    surface_sheets.append(name)
    fields = [e for e in ENTRIES if portal in PORTAL[e["portal"]] and any(AUD[c] in personas for c in e["aud"])]
    ws.append([screen])
    ws["A1"].font = Font(bold=True, size=14, color="1F3A8A")
    for label, val in (("Portal", portal.capitalize()), ("Where", where), ("Email section", SECTION_BY_SURFACE[name]),
                       ("Who can be emailed", ", ".join(personas)), ("Note", note), ("Fields on this sheet", len(fields))):
        ws.append([label, val])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.append([])
    head_row = ws.max_row + 1
    ws.append(FIELD_HEADS)
    for c in ws[head_row]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD; c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for e in fields:
        avail = [AUD[c] for c in e["aud"] if AUD[c] in personas]
        ws.append(["{{%s}}" % e["token"], e["label"], e["group"], e["desc"],
                   "All personas here" if len(avail) == len(personas) else ", ".join(avail),
                   e["prod"], e["example"] or "(typed by sender)", e["kind"].capitalize(), HOW[e["how"]],
                   "Existing" if norm(e["token"]) in legacy_tokens or e["token"] in ("firstName", "lastName", "email") else "New"])
        for c in ws[ws.max_row]:
            c.alignment = WRAP; c.border = BORDER
    for i, w_ in enumerate(FIELD_W, 1):
        ws.column_dimensions[get_column_letter(i)].width = w_
    ws.column_dimensions["B"].width = 40
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(FIELD_HEADS))}{ws.max_row}"
    index_rows.append([name, portal.capitalize(), screen, where, ", ".join(personas), len(fields), note])

ws = wb.create_sheet("Where Bulk Email Appears")
sheet(ws, ["Sheet", "Portal", "Screen", "Where", "Who can be emailed", "Fields", "Note"], index_rows, [28, 9, 42, 40, 60, 8, 90])
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=1).hyperlink = f"#'{ws.cell(row=r, column=1).value}'!A1"
    ws.cell(row=r, column=1).font = Font(color="1D4ED8", underline="single")

# 5 Portal Survey
ws = wb.create_sheet("Portal Survey")
survey = [
    ("Site", "Locations (list)", "/site-details/locations", "Location, Type, Location Address, Group, Location Coordinators, Status", "Location group: locationName, locationAddress, locationGroup, locationCoordinators, locationStatus, locationType"),
    ("Site", "Location detail > About", "/site-details/locations/{id}", "Location Name, Alias name, Phone, Parent, Address, Description, Associated Disciplines, Departments & Units, Onboarding Requirements, Groups", "locationName, locationPhone, locationParent, locationAddress, locationDisciplines, locationDepartment"),
    ("Site", "Location detail > Availability", "/site-details/locations/{id}/availability", "Availability Name, Availability ID, Created By, Discipline/Specialization, Experience Type, Total Slots, Pending Requests, Requested Slots, Created On, Availability Duration, Program Type, Status", "Availability group"),
    ("Site", "Location detail > Requests", "/site-details/locations/{id}/requests", "Request ID, School Details, Availability Name, Experience Type, Requested Slots, Request Pending Duration, Program Type, Requested Shifts, Requested Days of Week, Requested Date, Requested By, Requested Duration, Approved Info, Student Profile Shared, Partner Category, Status", "Slot Request group"),
    ("Site", "Location detail > Schedules", "/site-details/locations/{id}/schedules", "Schedule ID, Student, School Name, Availability, Schedule Duration, Discipline/Specialization, Shift, Days of Week, Preceptor, Clinical Instructor, Program Type, Schedule Status, Onboarding Status; All Schedules / Group / Individual; All / Upcoming / Current / Completed", "Schedule group; studentSchedules list"),
    ("Site", "Personnel (list)", "/site-details/personnel", "Name, Role & Access, Email, Designation, Last Login, Status; filters Role, Status, Designation, System Access", "Site Personnel group"),
    ("Site", "Personnel detail", "/site-details/personnel/{id}", "First name, Last name, National provider identifier, Email, Contact, Credentials, Designations, Address, Notes and documents, Account Audit, Access & Role, Associated Disciplines, Associated locations (with Coordinator role), Coordinator for groups", "staffName, credentials, npi, designation, accessRole, systemAccess, associatedLocations, associatedDisciplines"),
    ("Site", "School Partners (list)", "/site-details/partners/partner-programs", "School Name, Website, Address, Category, Contracts (Active / Expired), Program Contact (+N), Associated Discipline", "School / Partner group; Contract counts"),
    ("Site", "School Partner > About", "/partner-programs/{id}/about", "Name, Alias name, Category, Website, Associated Discipline, Address, Related Documents, Program Contacts (Name, Email, Phone)", "schoolName, schoolAlias, schoolCategory, schoolWebsite, schoolAddress, programContacts, programContactCount"),
    ("Site", "School Partner > Contracts", "/partner-programs/{id}/contract", "Total / Site-specific / Location-specific counts; Name and Number, Contract Type, Start Date, Expiration Date, Open Ended Contract, Template Type, Notes, Contract Documents", "Contract group"),
    ("Site", "School Partner tabs", "/partner-programs/{id}/…", "About, Contracts, Engagement Log, Requests, Schedules", "partnerSince, requests and schedules groups"),
    ("Site", "Availability (list / detail)", "/internships/availability, /internships/{id}", "Overview, Availability List, Map View, Reports; detail: Total slots, Total approved, Experience type, Disciplines, Duration, Program type, Student year, Start–End date, Shift, Days of the week, Number of slots, Location, Eligibility criteria, Onboarding requirements, Request Fields, Publish Options (Published to, Publish Status, Published On, Due On)", "Availability group"),
    ("Site", "Slot Requests (list)", "/school-requests/list", "Request ID, School Details, Availability Name, Experience Type, Requested Slots, Request Pending Duration (Days), Requested Location, Program Type, Requested Shifts, Requested Days of Week, Requested Date, Requested By (name, School/Site, email), Requested Duration, Approved Info (e.g. 1 (1 Revoked)), Student Profile Shared, Partner Category, Status (Approved / In Progress / Request Pending)", "Slot Request group"),
    ("Site", "Schedules (list)", "/assignments/list", "Schedule ID, Student (name, email), School Name, Availability, Location, Schedule Duration, Discipline/Specialization, Shift, Days of Week, Preceptor, Clinical Instructor, Program Type, Schedule Status (incl. cancellation reason), Onboarding Status; filters Schedule Type, Discipline, Schools, Start & End Date, Location Groups", "Schedule group"),
    ("Site", "Learner profile", "/learners-profile/{id}", "Name, Email, Flag for Hire, profile completion", "studentName, email"),
    ("Site", "Jobs (list)", "/jobs", "Title, Job Type, Location, Discipline, Status, Posted On, Applications, Under Review, Last Updated; tabs All / Draft / Active / Paused / Closed / Archived", "Candidate & Job group"),
    ("School", "Requests (list)", "/school/{id}/requests", "Site Name, Request ID, Location Name, Shift, Days of the Week, Discipline/Specialization, Requested Duration, Requested Slots, Date of Request, Confirmed Slots, Requested By (name, email), Experience Type, Request Pending Duration, Program Type, Status", "Slot Request group (confirmedSlots, requestedBy, requestedByEmail)"),
    ("School", "Schedules (list)", "/school/{id}/assignments/list", "Schedule ID, Student, Site & Location, Availability, Schedule Duration, Discipline/Specialization, Clinical Instructor, Program Type, Rotation & Course Details, PRISM Process Status, Schedule Status, Onboarding Status (incl. School Requirements); actions Confirm Students, Remind to Onboard; filter Student Login Status", "Schedule group (rotationCourseDetails, prismProcessStatus, schoolRequirementsStatus, studentLoginStatus)"),
    ("School", "Dashboard tasks", "/school/{id}/dashboard", "Student Scheduling and Confirmation Pending, Group Instructor Scheduling pending, messages", "openRequests, studentSchedules lists"),
    ("School", "Explore, Wishlist, Activity Dashboard, Reports", "/search-internships, /wishlist-management, /activity-dashboard, /reports", "Browsed for field names only; they are not Bulk Email surfaces today", "—"),
]
sheet(ws, ["Portal", "Screen", "Path", "What the screen shows", "Became these merge fields"], survey, [10, 30, 40, 100, 60])

# 6 Lists
ws = wb.create_sheet("Lists")
rows = [["{{#each %s}}" % e["token"], e["label"], ", ".join(audiences_of(e)), e["desc"], e["example"], e["insert"]] for e in ENTRIES if e["kind"] == "list"]
sheet(ws, ["Block", "Label", "Audiences", "What each line holds", "Example line", "Inserted as"], rows, [30, 30, 40, 50, 60, 60])

# 7 Notes
ws = wb.create_sheet("Notes for Development")
notes = [
    ("Schools with no program contacts", "When a school has no program contact details, send to the person who raised the request from the school side, using requestedByEmail. Product decision; the prototype does not model the fallback."),
    ("Sample fields", "Fields marked 'Sample (prototype)' exist in the live product but not in the prototype's records. In production, read them from the 'Where it lives' column."),
    ("HTML in live data", "Live records contain HTML and script text in name fields. Escape every merged value in HTML bodies (the prototype already does)."),
    ("Same token, two meanings", "{{preceptorName}} is the assigned preceptor for a student and the recipient's own name for a preceptor. {{department}} is the department for staff and for a job."),
    ("Mixed-persona sends", "When one message goes to several personas, each recipient's fields fill from their own record; a field that does not apply to a recipient holds that recipient back and is flagged."),
    ("Not covered", "Availability and Reports are not Bulk Email surfaces yet, so no section-level fields were added for them beyond the Availability group."),
]
sheet(ws, ["Topic", "Note"], [list(n) for n in notes], [34, 130])

ws = wb.create_sheet("Removed Fields")
rows = [["{{%s}}" % t, v[2], v[3], v[0], v[1]] for t, v in sorted(REMOVED.items(), key=lambda kv: (GROUP_ORDER.index(kv[1][3]), kv[1][0]))]
sheet(ws, ["Token", "Label", "Group", "Why removed", "Reason"], rows, [32, 34, 24, 24, 90])

order = ["Summary", "Where Bulk Email Appears"] + surface_sheets + ["All Merge Fields", "By Persona", "Removed Fields", "Portal Survey", "Lists", "Notes for Development"]
wb._sheets = [wb[n] for n in order]
wb.active = 0
out = ROOT / "Merge_Fields_Catalog.xlsx"
wb.save(out)
print("workbook ->", out.name)
